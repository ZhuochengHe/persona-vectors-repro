"""Build/update the reading workbook with stable IDs and three-way cell merging."""

from __future__ import annotations

import argparse
import csv
import json
import os
import shutil
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs/papers"
DEFAULT = Path("/mnt/d/umich/research/persona/persona_literature.xlsx")
HEADERS = ["ID", "Title", "Link", "Key words", "One sentence/topic", "Contributions",
           "Limitation", "Additional think", "Related paper", "Research direction",
           "Priority", "Institution / Venue", "Read status", "Evidence scope"]
WIDTHS = [7, 62, 34, 38, 56, 62, 62, 58, 18, 26, 16, 52, 16, 34]
USER_FIELDS = {"Additional think", "Read status"}


def records() -> list[tuple[str, dict]]:
    catalog = json.loads((DOCS / "citation_screening_sources/workbook_catalog.json").read_text())
    with (DOCS / "citation_screening_shortlist.csv").open(newline="") as handle:
        selected = {row["record_id"]: row for row in csv.DictReader(handle)}
    result = []
    for key, identifier, keywords, topic, related in catalog:
        values = dict.fromkeys(HEADERS, "")
        values.update(ID=identifier, **{"Key words": keywords, "One sentence/topic": topic,
                      "Related paper": ", ".join(map(str, related)), "Read status": "未标记"})
        if key == "seed":
            values.update(Title="Persona Vectors: Monitoring and Controlling Character Traits in Language Models",
                          Link="https://arxiv.org/abs/2507.21509",
                          Contributions="给出自动化对比提取流程；将persona方向用于部署监测和控制、微调漂移预测、预防性干预及训练数据筛选。",
                          Limitation="主要在有限开放模型与trait上验证；依赖提示诱发和自动评分，方向的语义纯度及跨情境泛化仍需检验。",
                          **{"Research direction": "基础论文／人格向量", "Priority": "基础", "Institution / Venue": "来源：原始论文；机构和会议信息本表未重新核实",
                             "Evidence scope": "本地原文方法及主要实验；非完整逐项审计"})
        elif key == "games":
            values.update(Title="Persona Vectors in Games: Measuring and Steering Strategies via Activation Vectors",
                          Link="https://arxiv.org/abs/2603.21398",
                          Contributions="将利他、宽恕及对他人预期的方向用于游戏；发现措辞与策略可分离，自身行为与他人预期的方向也不完全相同。",
                          Limitation="主要是给定游戏描述下的选择；不等于真实多轮博弈。部分评分依赖judge，论文与公开代码的题目及设置应分开核对。",
                          **{"Research direction": "基础应用／游戏与策略", "Priority": "基础", "Institution / Venue": "预印本；主会与署名机构本表未重新核实",
                             "Evidence scope": "论文方法、Figure 3及复现协议核查；非完整逐项审计"})
        else:
            row = selected[key]
            values.update(Title=row["title"], Link=row["paper_url"],
                          Contributions=row["main_contribution"], Limitation=row["limitation_or_reading_caution"],
                          **{"Research direction": row["shortlist_group"], "Priority": row["shortlist_tier"],
                             "Institution / Venue": row["verified_affiliations"] + "；" + row["venue"],
                             "Evidence scope": "官方摘要＋署名及正文贡献／讨论／局限片段；非全文审计"})
        result.append((key, values))
    ids = {row["ID"] for _, row in result}
    assert len(ids) == len(result) == 32
    for _, row in result:
        assert all(row[field] for field in ["Title", "Link", "Key words", "One sentence/topic"])
        assert all(int(n) in ids and int(n) != row["ID"] for n in row["Related paper"].split(", "))
    return result


def guide(book: Workbook) -> None:
    sheet = book.create_sheet("Guide")
    items = [
        ("用途", "Windows端高层阅读整理；研究过程、筛选和引用证据继续在仓库维护。"),
        ("目标位置", r"D:\umich\research\persona\persona_literature.xlsx"),
        ("WSL位置", str(DEFAULT)),
        ("内容", "32篇：2篇基础论文＋30篇精选；ID 1/2固定为Persona Vectors及游戏应用。"),
        ("ID", "稳定论文编号，不是行号；筛选、排序后Related paper仍按这个编号理解。"),
        ("Related paper", "主题、方法或结论相关的论文编号；不表示这些论文已确认互相引用。"),
        ("Additional think", "专留给你填写，初始化为空，更新脚本不会生成或覆盖此列。"),
        ("阅读状态", "未标记／待读／在读／已读／暂缓；更新时保留。"),
        ("优先级", "基础：最初两篇；A：先建立问题地图的10篇；B：按方向扩展的20篇。"),
        ("证据边界", "贡献是当前综合判断，Limitation包含作者局限与阅读核对点；未声称32篇已全文审计。"),
        ("更新", "先更新仓库文献源数据，再运行scripts/export_literature_workbook.py；按稳定ID合并。"),
        ("手动修改", "自动字段仅在当前内容仍等于上次导出值时更新；你改过的内容保留。已存在的行不自动删除。"),
        ("保管", "更新前关闭Excel；脚本会为已有文件保留备份。_ExportState是内部合并基线，不要删除。"),
        ("生成时间 UTC", datetime.now(timezone.utc).isoformat()),
    ]
    sheet.append(["说明", "内容"])
    for row in items:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 23
    sheet.column_dimensions["B"].width = 110
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 36
    sheet.freeze_panes = "B2"


def export(output: Path) -> None:
    if not output.parent.is_dir():
        raise SystemExit(f"Destination directory is unavailable: {output.parent}. Run in WSL or specify an existing staging directory; no substitute mount will be created.")
    if output.with_name("~$" + output.name).exists():
        raise SystemExit("Excel lock file exists; close the workbook before updating.")
    expected = records()
    fresh = not output.exists()
    if fresh:
        book = Workbook()
        sheet = book.active
        sheet.title = "Papers"
        sheet.append(HEADERS)
        guide(book)
        state = book.create_sheet("_ExportState")
        state.append(["Source key", "Stable ID", "Last generated values"])
        baseline = {}
    else:
        book = load_workbook(output)
        if not {"Papers", "_ExportState"}.issubset(book.sheetnames):
            raise SystemExit("Existing workbook lacks merge metadata; inspect and migrate it before updating.")
        sheet, state = book["Papers"], book["_ExportState"]
        if [cell.value for cell in sheet[1]] != HEADERS:
            raise SystemExit("Workbook headers changed; inspect before updating.")
        baseline = {key: (identifier, json.loads(values)) for key, identifier, values in state.iter_rows(min_row=2, values_only=True) if key}
    row_by_id = {}
    for row in sheet.iter_rows(min_row=2):
        if row[0].value is not None:
            if row[0].value in row_by_id:
                raise SystemExit("Duplicate paper IDs; inspect before updating.")
            row_by_id[row[0].value] = row[0].row
    next_baseline = dict(baseline)
    new_rows = []
    for key, values in expected:
        identifier = values["ID"]
        previous = baseline.get(key)
        if previous and previous[0] != identifier:
            raise SystemExit(f"Stable ID changed for {key}; inspect before updating.")
        if identifier in row_by_id and not previous:
            raise SystemExit(f"ID {identifier} is occupied without a matching source key.")
        if previous and identifier not in row_by_id:
            # Respect a user's deletion rather than silently recreating the row.
            continue
        new = identifier not in row_by_id
        if new:
            sheet.append([values[header] for header in HEADERS])
            row_number = sheet.max_row
            row_by_id[identifier] = row_number
            new_rows.append(row_number)
        else:
            row_number = row_by_id[identifier]
        for column, header in enumerate(HEADERS, 1):
            cell = sheet.cell(row_number, column)
            old = previous[1].get(header) if previous else None
            unchanged = cell.value == old or (cell.value is None and old == "")
            if new or (header not in USER_FIELDS and unchanged):
                if header == "Link" and not new and cell.hyperlink and cell.hyperlink.target != old:
                    continue
                cell.value = values[header]
                if header == "Link":
                    cell.hyperlink = values[header]
                    cell.font = Font(name="Calibri", size=11, color="0563C1", underline="single")
        next_baseline[key] = (identifier, values)
    if fresh:
        sheet.freeze_panes = "C2"
        sheet.sheet_view.showGridLines = False
        for column, width in enumerate(WIDTHS, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 34
        validation = DataValidation(type="list", formula1='"未标记,待读,在读,已读,暂缓"', allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add("M2:M10000")
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
    for number in new_rows:
        sheet.row_dimensions[number].height = 112
        for cell in sheet[number]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column != 3:
                cell.font = Font(name="Calibri", size=11)
        sheet.cell(number, 8).fill = PatternFill("solid", fgColor="FFF2CC")
    reference = f"A1:N{sheet.max_row}"
    if "LiteratureTable" in sheet.tables:
        sheet.tables["LiteratureTable"].ref = reference
        if sheet.tables["LiteratureTable"].autoFilter:
            sheet.tables["LiteratureTable"].autoFilter.ref = reference
    else:
        table = Table(displayName="LiteratureTable", ref=reference)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    state.delete_rows(2, max(state.max_row - 1, 0))
    for key, (identifier, values) in sorted(next_baseline.items(), key=lambda item: item[1][0]):
        state.append([key, identifier, json.dumps(values, ensure_ascii=False)])
    state.sheet_state = "veryHidden"
    book.active = book.sheetnames.index("Papers")
    temporary = output.with_name(output.stem + ".tmp.xlsx")
    try:
        book.save(temporary)
        if not fresh:
            stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
            shutil.copy2(output, output.with_name(output.stem + f".backup-{stamp}.xlsx"))
        os.replace(temporary, output)
    finally:
        book.close()
        temporary.unlink(missing_ok=True)
    print(f"Saved {output} ({len(expected)} catalog papers; user notes preserved).")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT)
    export(parser.parse_args().output)
